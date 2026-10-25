"""نقطهٔ ورود برنامه — FastAPI."""
from __future__ import annotations

import io
import json as _json
import time
from datetime import datetime
from pathlib import Path

from fastapi import Request, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import ai, catalog, engine, views
from .config import (APP_VERSION, APP_BUILD, DEFAULT_PERIOD, DEFAULT_YEAR,
                     MODEL_ANALYSIS, MODEL_CHAT, MODEL_DEEP, OLLAMA_HOST, PERIOD_LABEL)
from .db import current_version, get_db, new_version
from .seed import seed

app = FastAPI(title="داشبورد مدیریتی مجتمع خراسان", docs_url="/api/docs")
STATIC = Path(__file__).parent / "static"
class NoCacheStatic(StaticFiles):
    """بدون no-cache، مرورگر app.js قدیمی را نگه می‌دارد و نسخهٔ تازه دیده نمی‌شود."""
    async def get_response(self, path, scope):
        resp = await super().get_response(path, scope)
        if path.endswith((".js", ".css", ".html")):
            resp.headers["Cache-Control"] = "no-cache, must-revalidate"
        return resp


app.mount("/static", NoCacheStatic(directory=str(STATIC)), name="static")


@app.on_event("startup")
def _startup():
    info = seed()
    with get_db() as conn:
        engine.run(conn, DEFAULT_YEAR, DEFAULT_PERIOD)
    print("seed:", info)


# ── صفحات
@app.get("/")
def index():
    """HTML صفحه با شمارهٔ نسخه در آدرس فایل‌های استاتیک.
    بدون این، مرورگر app.js قدیمی را از حافظه می‌خواند و تغییرات دیده نمی‌شود."""
    html = (STATIC / "index.html").read_text(encoding="utf-8")
    html = html.replace("/static/styles.css", f"/static/styles.css?v={APP_VERSION}")
    html = html.replace("/static/app.js", f"/static/app.js?v={APP_VERSION}")
    html = html.replace("%%VERSION%%", APP_VERSION)
    return HTMLResponse(html, headers={"Cache-Control": "no-cache, must-revalidate"})


@app.get("/api/meta")
def meta():
    with get_db() as conn:
        branches = [r["name"] for r in conn.execute(
            "SELECT name FROM dim_branch WHERE active=1 ORDER BY rowid")]
        ver = current_version(conn)
        years = [r["y"] for r in conn.execute("SELECT DISTINCT year y FROM fact ORDER BY y")]
    return {
        "branches": branches, "years": years or [1405],
        "periods": [{"key": k, "label": v} for k, v in PERIOD_LABEL.items()],
        "default": {"year": DEFAULT_YEAR, "period": DEFAULT_PERIOD},
        "data_version": ver,
        "models": {"chat": MODEL_CHAT, "analysis": MODEL_ANALYSIS, "deep": MODEL_DEEP},
        "ollama_host": OLLAMA_HOST,
        "domains": [{"key": k, "label": v} for k, v in catalog.DOMAINS],
        "app_version": APP_VERSION,
        "app_build": APP_BUILD,
        "templates": {k: v[0] for k, v in catalog.TEMPLATES.items()},
    }


@app.get("/api/dashboard")
def dashboard(page: str = "overview", branch: str = "مجتمع",
              year: int = DEFAULT_YEAR, period: str = DEFAULT_PERIOD):
    if page not in views.PAGE_TITLES:
        raise HTTPException(404, "صفحه یافت نشد")
    with get_db() as conn:
        return views.build_page(conn, page, branch, year, period)


# ── هوش مصنوعی
@app.get("/api/ai/health")
def ai_health(request: Request):
    out = ai.health()
    # تشخیص اینکه رابط کاربری از رایانهٔ دیگری باز شده باشد
    # (مثلاً پیش‌نمایش ابری) — در آن حالت Ollama روی «همان» رایانه در دسترس نیست
    host_hdr = (request.headers.get("host") or "").split(":")[0].lower()
    local_names = {"localhost", "127.0.0.1", "0.0.0.0", "::1", ""}
    out["remote_backend"] = host_hdr not in local_names and not host_hdr.endswith(".local")
    out["host_header"] = host_hdr
    if not out["online"] and out["remote_backend"]:
        out["error_text"] = ("این صفحه از رایانهٔ دیگری باز شده و برنامهٔ سمت‌سرور "
                             "به Ollama روی رایانهٔ شما دسترسی ندارد. "
                             "داشبورد باید روی همان رایانه‌ای اجرا شود که Ollama روی آن است.")
    return out


@app.get("/api/ai/diag")
def ai_diag():
    """گزارش عیب‌یابی کامل — قابل مشاهده در خود داشبورد."""
    return ai.diag()


@app.get("/api/ai/selftest")
def ai_selftest(model: str | None = None):
    """آزمون گام‌به‌گام اتصال و پاسخ مدل."""
    return ai.selftest(model)


@app.get("/api/ai/ps")
def ai_ps():
    return {"loaded": ai.ps()}


class WarmReq(BaseModel):
    model: str | None = None


@app.post("/api/ai/warmup")
def ai_warmup(req: WarmReq):
    return ai.warmup(req.model)


class AnalyzeReq(BaseModel):
    page: str = "overview"
    branch: str = "مجتمع"
    year: int = DEFAULT_YEAR
    period: str = DEFAULT_PERIOD
    model: str | None = None
    force: bool = False


class ChatMsg(BaseModel):
    role: str
    content: str


class ChatReq(BaseModel):
    messages: list[ChatMsg]
    model: str | None = None


@app.post("/api/ai/chat")
def ai_chat(req: ChatReq):
    """گفت‌وگوی آزاد با مدل — برای آزمودن اتصال و پرسش‌های عمومی."""
    want = req.model or MODEL_CHAT
    used, subbed = ai.pick_model(want)
    msgs = [{"role": m.role, "content": m.content} for m in req.messages]
    t0 = time.time()
    try:
        raw = ai._chat(used, msgs, max_tokens=600, json_mode=False)
        return {"status": "ok", "source": "ollama", "model": used,
                "model_notice": (f"مدل «{want}» نصب نیست؛ با «{used}» پاسخ داده شد."
                                 if subbed else None),
                "reply": raw, "latency_ms": int((time.time() - t0) * 1000)}
    except Exception as e:  # noqa: BLE001
        return {"status": "error", "source": "none", "model": used,
                "reply": "", "error": f"{type(e).__name__}: {e}",
                "latency_ms": int((time.time() - t0) * 1000)}


@app.post("/api/ai/chat/stream")
def ai_chat_stream(req: ChatReq):
    """همان گفت‌وگو، ولی تدریجی (NDJSON) — کاربر می‌بیند مدل واقعاً کار می‌کند."""
    want = req.model or MODEL_CHAT
    used, _ = ai.pick_model(want)

    def gen():
        first = True
        for kind, val in ai._chat_stream(used, [{"role": m.role, "content": m.content}
                                                for m in req.messages]):
            if kind == "توکن":
                yield _json.dumps({"t": "tok", "v": val}, ensure_ascii=False) + "\n"
            elif kind == "خطا":
                yield _json.dumps({"t": "err", "v": val}, ensure_ascii=False) + "\n"
            elif kind == "نکته":
                yield _json.dumps({"t": "note", "v": val}, ensure_ascii=False) + "\n"
            else:
                yield _json.dumps({"t": "end", "v": {**val, "model_used": used}},
                                                ensure_ascii=False) + "\n"
            first = False

    return StreamingResponse(gen(), media_type="application/x-ndjson",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


class QuickReq(BaseModel):
    page: str = "flt"
    branch: str = "مجتمع"
    year: int = DEFAULT_YEAR
    period: str = DEFAULT_PERIOD
    model: str | None = None


@app.post("/api/ai/quick")
def ai_quick(req: QuickReq):
    """تحلیل کوتاه یک دامنه — مثلاً فقط ماشین‌آلات."""
    with get_db() as conn:
        data = views.build_page(conn, req.page, req.branch, req.year, req.period)
        return ai.quick(conn, f"q:{req.page}:{req.branch}:{req.year}:{req.period}",
                        f"{data['title']} — {req.branch}", data["bundle"],
                        model=req.model, focus=req.page)


@app.post("/api/ai/analyze")
def ai_analyze(req: AnalyzeReq):
    with get_db() as conn:
        data = views.build_page(conn, req.page, req.branch, req.year, req.period)
        res = ai.analyze(
            conn, f"{req.page}:{req.branch}:{req.year}:{req.period}",
            f"{data['title']} — {req.branch}", data["bundle"], data["alerts"],
            model=req.model, force=req.force,
        )
        return res


@app.get("/api/ai/history")
def ai_history(limit: int = 20):
    with get_db() as conn:
        rows = [dict(r) for r in conn.execute(
            "SELECT * FROM ai_log ORDER BY id DESC LIMIT ?", (limit,))]
    return {"items": rows}


class AskReq(BaseModel):
    question: str
    page: str = "overview"
    branch: str = "خودکار"
    year: int = DEFAULT_YEAR
    period: str = DEFAULT_PERIOD
    model: str | None = None


@app.post("/api/ai/ask")
def ai_ask(req: AskReq):
    """پرسش آزاد از دادهٔ واقعی — «نظرت در مورد شعبهٔ نیشابور چیه؟»"""
    t0 = time.time()
    with get_db() as conn:
        try:
            res = ai.ask(conn, req.question, branch=req.branch,
                         year=req.year, period=req.period, model=req.model)
            return {"status": "ok", "source": "ollama", **res,
                    "latency_ms": int((time.time() - t0) * 1000)}
        except Exception as e:  # noqa: BLE001
            return {"status": "error", "source": "none", "answer": "",
                    "error": f"{type(e).__name__}: {e}",
                    "latency_ms": int((time.time() - t0) * 1000)}


@app.post("/api/ai/ask/stream")
def ai_ask_stream(req: AskReq):
    """همان پرسش، تدریجی — کاربر می‌بیند مدل با داده کار می‌کند."""
    t0 = time.time()

    def gen():
        with get_db() as conn:
            try:
                for kind, val in ai.ask_stream(conn, req.question, branch=req.branch,
                                               year=req.year, period=req.period,
                                               model=req.model):
                    if kind == "توکن":
                        yield _json.dumps({"t": "tok", "v": val}, ensure_ascii=False) + "\n"
                    elif kind == "نکته":
                        yield _json.dumps({"t": "note", "v": val}, ensure_ascii=False) + "\n"
                    elif kind == "خطا":
                        yield _json.dumps({"t": "err", "v": val}, ensure_ascii=False) + "\n"
                    else:
                        yield _json.dumps({"t": "end", "v": {**val,
                                            "latency_ms": int((time.time() - t0) * 1000)}},
                                          ensure_ascii=False) + "\n"
            except Exception as e:  # noqa: BLE001
                yield _json.dumps({"t": "err", "v": f"{type(e).__name__}: {e}"},
                                  ensure_ascii=False) + "\n"

    return StreamingResponse(gen(), media_type="application/x-ndjson",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/ai/branches")
def ai_branches():
    """فهرست شعبه‌ها برای منوی صفحهٔ گفت‌وگو."""
    return {"items": ["خودکار"] + ai.BRANCHES}

@app.get("/api/data/quality")
def data_quality(year: int = DEFAULT_YEAR, period: str = DEFAULT_PERIOD):
    with get_db() as conn:
        return views.build_page(conn, "data", "مجتمع", year, period)


@app.get("/api/template/{domain}")
def template(domain: str, year: int = DEFAULT_YEAR, period: str = DEFAULT_PERIOD):
    if domain not in catalog.TEMPLATES:
        raise HTTPException(404, "قالب یافت نشد")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    label, codes = catalog.TEMPLATES[domain]
    wb = Workbook()
    ws = wb.active
    ws.title = "داده"
    ws.sheet_view.rightToLeft = True

    head = ["شعبه", "سال", "دوره", "کد قلم", "عنوان قلم", "ماه", "مقدار", "سناریو"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF", name="Vazirmatn")
        cell.fill = PatternFill("solid", fgColor="0E2A47")
        cell.alignment = Alignment(horizontal="center")

    branches = ["مجتمع"] + [b for b in [
        "شهید کاوه", "مطهری", "سرخس", "لطف آباد", "نیشابور", "ماهیرود",
        "دوغارون", "بجنورد", "بیرجند", "باجگیران"]]
    r = 2
    for code in codes:
        title = catalog.PAY_ROWS.get(code) or catalog.CODE_LABELS.get(code, code)
        ws.append([branches[0], year, period, code, title, 0, None, "ACT"])
        for i in range(1, 7):
            ws.append([branches[0], year, period, code, title, i, None, "ACT"])
        r += 7
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["E"].width = 34
    for col in "BCDFGH":
        ws.column_dimensions[col].width = 12

    dv_b = DataValidation(type="list", formula1='"' + ",".join(branches) + '"', allow_blank=False)
    ws.add_data_validation(dv_b)
    dv_b.add(f"A2:A{r}")
    dv_s = DataValidation(type="list", formula1='"ACT,BUD,LY"', allow_blank=False)
    ws.add_data_validation(dv_s)
    dv_s.add(f"H2:H{r}")

    ws2 = wb.create_sheet("راهنما")
    ws2.sheet_view.rightToLeft = True
    guide = [
        [f"قالب ورود داده — {label}"],
        [""],
        ["۱. هر ردیف یک قلم داده است: شعبه، سال، دوره، کد قلم، ماه، مقدار."],
        ["۲. ستون «مقدار» تنها ستونی است که پر می‌کنید. ماه = ۰ یعنی جمع دوره."],
        ["۳. سناریو: ACT عملکرد واقعی، BUD بودجه، LY دوره مشابه سال قبل."],
        ["۴. واحدها: مبالغ به میلیون ریال، تناژ به تن، مساحت به متر مربع."],
        ["۵. پس از پر کردن، فایل را در صفحهٔ «ورود داده» بارگذاری کنید."],
        ["۶. ردیف‌های قفل‌شده (شعبه/کد/ماه) را تغییر ندهید تا کپی-پیست بازه‌ای کار کند."],
        [""],
        ["کدهای این قالب:"],
    ]
    for row in guide:
        ws2.append(row)
    ws2.append(["کد", "عنوان"])
    for code in codes:
        ws2.append([code, catalog.PAY_ROWS.get(code) or catalog.CODE_LABELS.get(code, code)])
    ws2.column_dimensions["A"].width = 90
    ws2.column_dimensions["B"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"template_{domain}_{year}_{period}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'})


@app.post("/api/data/import")
async def import_excel(file: UploadFile = File(...)):
    from openpyxl import load_workbook

    raw = await file.read()
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"فایل خوانده نشد: {e}") from e

    ws = wb["داده"] if "داده" in wb.sheetnames else wb.active
    rows, errors, warnings = [], [], []
    unknown_codes = set()
    n = 0
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v is None for v in row):
            continue
        branch, year, period, code, title, month, value, scenario = (list(row) + [None] * 8)[:8]
        if value is None or value == "":
            continue
        n += 1
        if not branch or not code:
            errors.append({"row": i, "msg": "شعبه یا کد قلم خالی است"})
            continue
        try:
            value = float(value)
        except (TypeError, ValueError):
            errors.append({"row": i, "msg": f"مقدار غیرعددی: {value}"})
            continue
        if value < 0 and not str(code).startswith(("FIN", "PAY")):
            errors.append({"row": i, "msg": f"مقدار منفی غیرمجاز برای {code}"})
            continue
        rows.append((str(branch).strip(), int(year or DEFAULT_YEAR), str(period or DEFAULT_PERIOD),
                     str(code).strip(), int(month or 0), str(scenario or "ACT"), value))

    if errors:
        return JSONResponse({"status": "error", "errors": errors[:40],
                             "error_count": len(errors), "accepted": 0}, status_code=422)

    with get_db() as conn:
        before = conn.execute("SELECT COUNT(*) c FROM fact").fetchone()["c"]
        conn.executemany(
            "INSERT OR REPLACE INTO fact(branch,year,period,code,month,scenario,value,version)"
            " VALUES(?,?,?,?,?,?,?,1)", rows)
        after = conn.execute("SELECT COUNT(*) c FROM fact").fetchone()["c"]
        vid = new_version(conn, "بارگذاری اکسل", f"{n} ردیف از فایل {file.filename}")
        conn.execute("UPDATE fact SET version=?", (vid,))
        engine.run(conn, DEFAULT_YEAR, DEFAULT_PERIOD)

    return {"status": "ok", "accepted": len(rows), "new_facts": after - before,
            "errors": [], "warnings": warnings, "version": vid}


class FactReq(BaseModel):
    branch: str
    code: str
    year: int = DEFAULT_YEAR
    period: str = DEFAULT_PERIOD
    month: int = 0
    scenario: str = "ACT"
    value: float


@app.post("/api/data/fact")
def save_fact(req: FactReq):
    with get_db() as conn:
        conn.execute(
            "INSERT OR REPLACE INTO fact(branch,year,period,code,month,scenario,value,version)"
            " VALUES(?,?,?,?,?,?,?,1)",
            (req.branch, req.year, req.period, req.code, req.month, req.scenario, req.value))
        vid = new_version(conn, "ورود دستی", f"{req.branch} / {req.code}")
        conn.execute("UPDATE fact SET version=?", (vid,))
        engine.run(conn, req.year, req.period)
    return {"status": "ok", "version": vid}


@app.get("/api/export/kpi.xlsx")
def export_kpi(year: int = DEFAULT_YEAR, period: str = DEFAULT_PERIOD):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "شاخص‌ها"
    ws.sheet_view.rightToLeft = True
    head = ["شاخص", "کد", "شعبه", "عملکرد", "بودجه", "دوره مشابه", "واحد", "فرمول"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E2A47")
        cell.alignment = Alignment(horizontal="center")
    meta = {r[0]: r for r in catalog.KPI_CATALOG}
    with get_db() as conn:
        for r in conn.execute("SELECT * FROM kpi WHERE year=? AND period=? ORDER BY branch, code",
                              (year, period)):
            m = meta.get(r["code"], ("", r["code"], "", ""))
            ws.append([m[1], r["code"], r["branch"], r["act"], r["bud"], r["ly"], m[3], m[5]])
    for col, w in zip("ABCDEFGH", (34, 16, 16, 16, 16, 16, 12, 50)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="kpi_{year}_{period}.xlsx"'})


@app.post("/api/admin/recompute")
def recompute(year: int = DEFAULT_YEAR, period: str = DEFAULT_PERIOD):
    with get_db() as conn:
        return engine.run(conn, year, period)
